"""Weekly run: pull new MCA filings, drop ones already delivered, enrich, write files.

Usage:  python -m mca_leads.run [--days 10] [--states CO,CT] [--no-enrich] [--out output]
"""
from __future__ import annotations

import argparse
import csv
import logging
import os
import smtplib
from datetime import date, timedelta
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .enrich import enrich
from .funders import FunderList
from .sources import colorado, connecticut, inbox, oregon, rhode_island

COLUMNS = ["status", "business_name", "email", "email_status", "website", "website_match", "owner_name",
           "address", "city", "state", "zip", "funder_or_agent", "mca_evidence", "filing_date", "state_filed",
           "file_id", "other_names", "email_source", "excluded_reason", "source"]
ATTRIBUTION = ("Sources: public UCC records of the Colorado Secretary of State (data.colorado.gov), Connecticut "
               "Secretary of the State (data.ct.gov), Oregon Secretary of State (data.oregon.gov) and Rhode Island "
               "Department of State, provided as-is; business websites/emails partly from Overture Maps Places "
               "(Contains Overture Maps data, CDLA-Permissive-2.0 / Apache-2.0). A UCC filing does not prove a balance is owed.")
log = logging.getLogger("mca_leads")


def _read_set(path: Path, col: str) -> set[str]:
    if not path.exists():
        return set()
    with path.open() as f:
        return {row[col].strip().lower() for row in csv.DictReader(f) if row.get(col)}


def save_seen(path: Path, keys: set[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["key"])
        w.writerows([k] for k in sorted(keys))


def mark(leads: list[dict], suppressed: set[str] | None = None) -> list[dict]:
    """Set status: Stacked / Lead / Possible (agent-only) / Suppressed / Excluded."""
    suppressed = suppressed or set()
    counts: dict[str, int] = {}
    for l in leads:
        k = l["business_name"].lower().strip()
        counts[k] = counts.get(k, 0) + 1
    for l in leads:
        email = (l.get("email") or "").lower()
        domain = email.split("@")[-1] if email else ""
        if l.get("excluded_reason"):
            l["status"] = "Excluded"
        elif email and (email in suppressed or domain in suppressed):
            l["status"], l["email"] = "Suppressed", ""
        elif counts[l["business_name"].lower().strip()] > 1:
            l["status"] = "Stacked"
        elif "possible" in (l.get("mca_evidence") or ""):
            l["status"] = "Possible"
        else:
            l["status"] = "Lead"
    order = {"Stacked": 0, "Lead": 1, "Possible": 2, "Suppressed": 3, "Excluded": 4}
    return sorted(leads, key=lambda l: (order[l["status"]], not l.get("email"), l["business_name"].lower()))


def sendable(l: dict) -> bool:
    return bool(l.get("email")) and l["status"] in ("Stacked", "Lead", "Possible") and \
        l.get("email_status", "not checked") in ("ok", "mx ok", "not checked")


def write_files(leads: list[dict], out: Path, stamp: str) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    csv_path, xlsx_path = out / f"mca_leads_{stamp}.csv", out / f"mca_leads_{stamp}.xlsx"
    with csv_path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(leads)
    wb = Workbook()
    sheets = [("Ready to email", [l for l in leads if sendable(l)]), ("All leads", leads)]
    for i, (title, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(COLUMNS)
        for c in ws[1]:
            c.font = Font(bold=True)
        for l in rows:
            ws.append([l.get(c, "") for c in COLUMNS])
        ws.freeze_panes = "C2"
        for j, c in enumerate(COLUMNS, 1):
            wide = c in ("business_name", "email", "website", "funder_or_agent", "mca_evidence", "address")
            ws.column_dimensions[get_column_letter(j)].width = 30 if wide else 14
    note = wb.create_sheet("About")
    note.append([ATTRIBUTION])
    note.append(["Status: Stacked = 2+ MCA filings this run; Lead = MCA evidence; Possible = filed by an agent "
                 "used by MCA funders and banks, check before sending; Suppressed = on your opt-out list."])
    wb.save(xlsx_path)
    return [csv_path, xlsx_path]


def email_files(paths: list[Path], summary: str) -> None:
    host, to = os.getenv("SMTP_HOST"), os.getenv("REPORT_TO")
    if not (host and to):
        return
    msg = EmailMessage()
    msg["Subject"], msg["From"], msg["To"] = "Weekly MCA lead file", os.getenv("SMTP_USER", to), to
    msg.set_content(summary + "\n\n" + ATTRIBUTION)
    for p in paths:
        msg.add_attachment(p.read_bytes(), maintype="application", subtype="octet-stream", filename=p.name)
    with smtplib.SMTP(host, int(os.getenv("SMTP_PORT") or 587)) as s:
        s.starttls()
        s.login(os.getenv("SMTP_USER", ""), os.getenv("SMTP_PASSWORD", ""))
        s.send_message(msg)
    log.info("Emailed files to %s", to)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--days", type=int, default=10, help="look-back window; overlap is removed by history")
    ap.add_argument("--states", default="CO,CT,RI,OR")
    ap.add_argument("--no-enrich", action="store_true")
    ap.add_argument("--out", default="output")
    ap.add_argument("--data", default="data")
    ap.add_argument("--inbox", default="inbox")
    a = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    data = Path(a.data)
    end = date.today()
    start = end - timedelta(days=a.days)
    states = [s.strip().upper() for s in a.states.split(",")]
    funders = FunderList(data / "funders.csv")

    leads: list[dict] = []
    if "CO" in states:                       # run first: it teaches the funder list
        try:
            co = colorado.fetch(start, end)
        except Exception as e:
            log.error("CO source failed: %s", e)
            co = []
        learned = funders.learn([l["funder_or_agent"] for l in co])
        log.info("Funder list: %d new names learned from Colorado", learned)
        funders.save()
        leads += co
    for code, fn, lookback in (("CT", connecticut.fetch, a.days), ("RI", rhode_island.fetch, a.days),
                               ("OR", oregon.fetch, 62)):      # Oregon publishes monthly, so look back further
        if code in states:
            try:
                leads += fn(end - timedelta(days=lookback), end, funders=funders)
            except Exception as e:        # one state failing must not stop the others
                log.error("%s source failed: %s", code, e)
    leads += inbox.fetch(start, end, folder=a.inbox, funders=funders)

    seen = _read_set(data / "seen.csv", "key")
    new = [l for l in leads if f"{l['state_filed']}:{l['file_id']}".lower() not in seen]
    log.info("%d MCA filings found, %d new since last run", len(leads), len(new))
    if not a.no_enrich:
        new = enrich(new)
    new = mark(new, _read_set(data / "suppression.csv", "email_or_domain"))
    paths = write_files(new, Path(a.out), end.isoformat())
    save_seen(data / "seen.csv", seen | {f"{l['state_filed']}:{l['file_id']}".lower() for l in new})

    n_lead = sum(l["status"] in ("Stacked", "Lead", "Possible") for l in new)
    summary = (f"MCA leads {start}..{end}: {n_lead} new merchants "
               f"({sum(l['status'] == 'Stacked' for l in new)} stacked, "
               f"{sum(l['status'] == 'Possible' for l in new)} possible), {sum(sendable(l) for l in new)} ready to email.")
    log.info(summary)
    email_files(paths, summary)
    if os.getenv("GITHUB_STEP_SUMMARY"):
        Path(os.environ["GITHUB_STEP_SUMMARY"]).write_text(summary + "\n")


if __name__ == "__main__":
    main()
